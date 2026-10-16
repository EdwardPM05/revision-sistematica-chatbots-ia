"""Construye los anexos del artículo: 50 artículos seleccionados con su evidencia.

Genera:
  - Documentos/Anexo_50_articulos_seleccionados.xlsx (presentación en Excel)
  - Sección "Anexos" al final de Documentos/Articulo_revision_sistematica_chatbots_IA.docx

Fuentes de verdad:
  - Datos/incluidos_full.json (76 estudios incluidos, extracción y calidad)
  - Documentos/PRISMA_identificacion_consolidado.xlsx (hojas Para_cribado_1171 y Texto_completo)

Uso (desde la raíz del repositorio):
  python3 Scripts/build_anexos.py
Se puede re-ejecutar: reemplaza los anexos previos del .docx.
"""
import copy
import html
import json
import re
from collections import Counter
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATOS = ROOT / "Datos"
DOCS = ROOT / "Documentos"
ARTICULO = DOCS / "Articulo_revision_sistematica_chatbots_IA.docx"
PRISMA_XLSX = DOCS / "PRISMA_identificacion_consolidado.xlsx"
SALIDA_XLSX = DOCS / "Anexo_50_articulos_seleccionados.xlsx"

N_SELECCION = 50
ANIO_CORTE = 2026
# Estudios incluidos que ya se citan en la Introducción del artículo (ID de registro -> [n])
CITADOS_EN_ARTICULO = {"640": 4, "909": 15, "902": 17, "900": 20, "929": 22}

SECTOR = {
    "educacion": "Educación", "salud": "Salud", "salud_mental": "Salud mental",
    "gobierno_publico": "Gobierno / sector público", "banca_finanzas": "Banca y finanzas",
    "industria_empresa": "Industria y empresa", "turismo_hoteleria": "Turismo y hotelería",
    "atencion_cliente_comercio": "Atención al cliente y comercio",
    "general_multisector": "General / multisectorial", "otro": "Otro",
}
TIPO = {
    "encuesta": "Encuesta", "revision_sistematica": "Revisión sistemática",
    "evaluacion_sistema": "Evaluación de sistema", "cuasiexperimental": "Cuasiexperimental",
    "mixto": "Mixto", "cualitativo": "Cualitativo", "bibliometrico": "Bibliométrico",
    "experimental": "Experimental", "metaanalisis": "Metaanálisis",
    "estudio_caso": "Estudio de caso", "revision_alcance": "Revisión de alcance",
}
MADUREZ = {
    "no_aplica": "No aplica", "uso_comercial_masivo": "Uso comercial masivo",
    "piloto": "Piloto", "implementacion_real": "Implementación real", "prototipo": "Prototipo",
}


def etiqueta(mapa, valor):
    return mapa.get(valor, valor.replace("_", " ").capitalize())


# ---------------------------------------------------------------- datos
def cargar():
    estudios = json.loads((DATOS / "incluidos_full.json").read_text(encoding="utf-8"))
    wb = load_workbook(PRISMA_XLSX, read_only=True)
    cribado = {str(r[0]): r for r in wb["Para_cribado_1171"].iter_rows(min_row=2, values_only=True)}
    texto = {str(r[0]): r for r in wb["Texto_completo"].iter_rows(min_row=2, values_only=True)}
    resumen = [r for r in wb["Resumen_PRISMA"].iter_rows(min_row=2, values_only=True) if r[1] is not None]
    for e in estudios:
        c, t = cribado[e["id"]], texto[e["id"]]
        e["base"] = c[1]
        e["titulo"] = html.unescape(c[2] or t[1] or "").strip()
        e["autores_bd"] = c[3]
        e["anio"] = int(c[4] or t[3] or re.findall(r"20\d\d", e["ref_ieee"])[-1])
        e["revista"] = html.unescape(c[5] or "")
        e["citas"] = int(c[6]) if str(c[6]).isdigit() else 0
        m = re.search(r"10\.\d{4,9}/[^\s,;]+", f"{c[7] or ''} {e['ref_ieee']}")
        e["doi"] = m.group(0).rstrip(".") if m else ""
        e["url_doi"] = f"https://doi.org/{e['doi']}" if e["doi"] else ""
        url = re.match(r"https?://\S+", e["fuente_texto"] or "")
        e["url_texto"] = url.group(0) if url else ""
        e["decision_tc"] = t[6]
        e["citas_anio"] = e["citas"] / (ANIO_CORTE - e["anio"] + 1)
        e["autor_anio"] = f"{autor_corto(e['ref_ieee'])} ({e['anio']})"
    return estudios, resumen


def autor_corto(ref):
    """Primer apellido + 'et al.' si hay 3 o más autores (regla IEEE de la UA)."""
    autores = ref.split(', "')[0]
    nombres = [a for a in re.split(r",\s*(?:and\s+|y\s+)?|\s+and\s+|\s+y\s+", autores) if a.strip()]
    nombres = [n for n in nombres if n.strip() not in ("et al.", "et al")]
    apellido = lambda n: n.strip().split(" ")[-1]
    if "et al" in autores or len(nombres) >= 3:
        return f"{apellido(nombres[0])} et al."
    if len(nombres) == 2:
        return f"{apellido(nombres[0])} y {apellido(nombres[1])}"
    return apellido(nombres[0])


def seleccionar(estudios):
    """Orden de prioridad: citado en el artículo > puntaje de calidad > LATAM > citas por año."""
    clave = lambda e: (e["id"] in CITADOS_EN_ARTICULO, int(e["puntaje"]), e["latam"] == "si", e["citas_anio"])
    orden = sorted(estudios, key=clave, reverse=True)
    sel = orden[:N_SELECCION]
    # numeración final por sector y año para facilitar la lectura
    sel.sort(key=lambda e: (etiqueta(SECTOR, e["sector"]), -e["anio"], e["autor_anio"]))
    for n, e in enumerate(sel, 1):
        e["n"] = n
        if e["id"] in CITADOS_EN_ARTICULO:
            e["motivo"] = f"Citado en la Introducción como [{CITADOS_EN_ARTICULO[e['id']]}]"
        else:
            partes = [f"Calidad {e['puntaje']}/5"]
            if e["latam"] == "si":
                partes.append("evidencia latinoamericana")
            partes.append(f"{e['citas_anio']:.1f} citas/año")
            e["motivo"] = "; ".join(partes)
    return sel


# ---------------------------------------------------------------- Excel
AZUL = PatternFill("solid", fgColor="1F3864")
GRIS = PatternFill("solid", fgColor="F2F2F2")
FINO = Side(style="thin", color="BFBFBF")
BORDE = Border(top=FINO, bottom=FINO, left=FINO, right=FINO)


def encabezado(ws, fila, columnas, anchos):
    for i, (col, ancho) in enumerate(zip(columnas, anchos), 1):
        c = ws.cell(row=fila, column=i, value=col)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = AZUL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[fila].height = 32


def celda(ws, fila, col, valor, enlace=None, centro=False):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = Font(name="Arial", size=9, color="0563C1" if enlace else "000000", underline="single" if enlace else None)
    c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if centro else "left")
    c.border = BORDE
    if enlace:
        c.hyperlink = enlace
    if fila % 2 == 1:
        c.fill = GRIS
    return c


def construir_excel(sel, resumen, total_incluidos):
    wb = Workbook()

    # Hoja 1: Presentación
    ws = wb.active
    ws.title = "Presentacion"
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    filas = [
        ("CHATBOTS BASADOS EN INTELIGENCIA ARTIFICIAL: UNA REVISIÓN SISTEMÁTICA DE LA LITERATURA 2021-2025", None, None),
        ("Anexo: 50 artículos seleccionados y evidencia de su selección", None, None),
        ("Autor: PITTMAN MEDINA, Edward Miguel · Asesor: VIDAL Rischmoller, Julio Cesar", None, None),
        ("Universidad Autónoma del Perú · Curso: Desarrollo de Investigación", None, None),
        (None, None, None),
        ("FLUJO PRISMA 2020", "N", "Nota"),
    ]
    filas += [(r[0], r[1], r[2]) for r in resumen]
    filas += [("Artículos seleccionados para este anexo", N_SELECCION,
               f"Subconjunto de los {total_incluidos} incluidos; criterio de selección abajo")]
    filas += [
        (None, None, None),
        ("CRITERIO DE SELECCIÓN DE LOS 50 ARTÍCULOS", None, None),
        ("Se ordenaron los estudios incluidos con la siguiente prioridad y se tomaron los 50 primeros:", None, None),
        ("1. Estudio citado en la Introducción del artículo", None, None),
        ("2. Mayor puntaje de calidad metodológica (escala 0-5)", None, None),
        ("3. Evidencia de Latinoamérica (pertinencia para el contexto peruano)", None, None),
        ("4. Mayor número de citas por año desde su publicación", None, None),
        (None, None, None),
        ("CONTENIDO DEL LIBRO", None, None),
        ("Articulos_50", None, "Matriz de extracción: autor, año, título, revista, base, país, sector, tipo de estudio, muestra, chatbot, madurez, hallazgos, barreras, limitaciones, referencia IEEE y DOI"),
        ("Evidencia", None, "Trazabilidad de cada artículo: ID de registro en PRISMA_identificacion_consolidado.xlsx, base de datos, decisión a texto completo, enlace al texto completo consultado, criterios de calidad C1-C5, puntaje, citas y motivo de selección"),
        ("Resumen", None, "Distribución de los 50 artículos por sector, tipo de estudio, año, base de datos y región"),
    ]
    for i, (a, b, c) in enumerate(filas, 1):
        for j, v in enumerate((a, b, c), 1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if a and (a.isupper() or i <= 2):
            ws.cell(row=i, column=1).font = Font(name="Arial", size=11 if i > 2 else 12, bold=True, color="1F3864")
    for i, (a, *_ ) in enumerate(filas, 1):
        if a in ("FLUJO PRISMA 2020",):
            for j in range(1, 4):
                ws.cell(row=i, column=j).font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                ws.cell(row=i, column=j).fill = AZUL

    # Hoja 2: Matriz de los 50 artículos
    ws = wb.create_sheet("Articulos_50")
    cols = ["N°", "Autor(es) y año", "Título", "Revista / fuente", "Base de datos", "País", "Sector",
            "Tipo de estudio", "Muestra", "Chatbot evaluado", "Madurez", "Principales hallazgos",
            "Barreras / riesgos", "Limitaciones", "LATAM", "Calidad (0-5)", "Referencia IEEE", "DOI"]
    anchos = [5, 20, 45, 25, 14, 14, 16, 16, 25, 28, 16, 55, 40, 40, 7, 9, 60, 30]
    encabezado(ws, 1, cols, anchos)
    for f, e in enumerate(sel, 2):
        valores = [e["n"], e["autor_anio"], e["titulo"], e["revista"], e["base"], e["pais"],
                   etiqueta(SECTOR, e["sector"]), etiqueta(TIPO, e["tipo_estudio"]), e["muestra"],
                   e["chatbot"], etiqueta(MADUREZ, e["madurez"]), e["hallazgo"], e["barreras"],
                   e["limitaciones"], "Sí" if e["latam"] == "si" else "No", int(e["puntaje"]), e["ref_ieee"]]
        for c, v in enumerate(valores, 1):
            celda(ws, f, c, v, centro=c in (1, 15, 16))
        celda(ws, f, 18, e["doi"] or e["url_texto"], enlace=e["url_doi"] or e["url_texto"])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(sel) + 1}"

    # Hoja 3: Evidencia
    ws = wb.create_sheet("Evidencia")
    cols = ["N°", "ID registro", "Autor(es) y año", "Base de datos", "Decisión título/resumen",
            "Decisión texto completo", "Texto completo recuperado", "Enlace al texto completo consultado",
            "Enlace DOI", "C1", "C2", "C3", "C4", "C5", "Puntaje (0-5)", "Citado por",
            "Nota de la evaluación", "Motivo de selección"]
    anchos = [5, 10, 20, 14, 14, 14, 12, 45, 35, 5, 5, 5, 5, 5, 9, 9, 50, 34]
    encabezado(ws, 1, cols, anchos)
    for f, e in enumerate(sel, 2):
        valores = [e["n"], int(e["id"]), e["autor_anio"], e["base"], "Pasa a texto completo",
                   e["decision_tc"].capitalize(), "Sí" if e["recuperado"] == "si" else "No"]
        for c, v in enumerate(valores, 1):
            celda(ws, f, c, v, centro=c in (1, 2, 7))
        celda(ws, f, 8, e["url_texto"], enlace=e["url_texto"] or None)
        celda(ws, f, 9, e["url_doi"], enlace=e["url_doi"] or None)
        for k, q in enumerate(("q1", "q2", "q3", "q4", "q5")):
            celda(ws, f, 10 + k, int(e[q]), centro=True)
        celda(ws, f, 15, int(e["puntaje"]), centro=True)
        celda(ws, f, 16, e["citas"], centro=True)
        celda(ws, f, 17, e.get("nota", ""))
        celda(ws, f, 18, e["motivo"])
    nota = len(sel) + 3
    ws.cell(row=nota, column=1, value=(
        "Nota. ID registro = identificador del estudio en las hojas Para_cribado_1171 y Texto_completo de "
        "PRISMA_identificacion_consolidado.xlsx. C1-C5 = criterios de calidad metodológica aplicados a texto "
        "completo (1 = cumple, 0 = no cumple); se excluyeron los estudios con puntaje < 4 (código FT5). "
        "Citado por = número de citas registrado en la base de datos al momento de la búsqueda.")
    ).font = Font(name="Arial", size=9, italic=True)
    ws.merge_cells(start_row=nota, start_column=1, end_row=nota, end_column=12)
    ws.row_dimensions[nota].height = 40
    ws.cell(row=nota, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(sel) + 1}"

    # Hoja 4: Resumen
    ws = wb.create_sheet("Resumen")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    fila = 1
    bloques = [
        ("Sector", Counter(etiqueta(SECTOR, e["sector"]) for e in sel)),
        ("Tipo de estudio", Counter(etiqueta(TIPO, e["tipo_estudio"]) for e in sel)),
        ("Año de publicación", Counter(str(e["anio"]) for e in sel)),
        ("Base de datos", Counter(e["base"] for e in sel)),
        ("Madurez del chatbot", Counter(etiqueta(MADUREZ, e["madurez"]) for e in sel)),
        ("Evidencia latinoamericana", Counter("Sí" if e["latam"] == "si" else "No" for e in sel)),
        ("Puntaje de calidad", Counter(f"{e['puntaje']}/5" for e in sel)),
    ]
    for titulo, conteo in bloques:
        encabezado(ws, fila, [titulo, "n", "%"], [34, 10, 10])
        fila += 1
        items = sorted(conteo.items()) if titulo == "Año de publicación" else conteo.most_common()
        for k, v in items:
            celda(ws, fila, 1, k)
            celda(ws, fila, 2, v, centro=True)
            celda(ws, fila, 3, f"=B{fila}/{len(sel)}", centro=True).number_format = "0.0%"
            fila += 1
        celda(ws, fila, 1, "Total").font = Font(name="Arial", size=9, bold=True)
        celda(ws, fila, 2, f"=SUM(B{fila - len(items)}:B{fila - 1})", centro=True).font = Font(name="Arial", size=9, bold=True)
        fila += 2

    wb.save(SALIDA_XLSX)


# ---------------------------------------------------------------- Word
def run(p, texto, b=False, i=False, sz=22, color=None, u=False):
    r = p.add_run(texto)
    r.font.name = "Arial"
    r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    r.bold, r.italic = b, i
    s = OxmlElement("w:sz")  # sz en medios puntos
    s.set(qn("w:val"), str(sz))
    r._r.get_or_add_rPr().append(s)
    return r


def espaciado(p, antes=0, despues=0, linea=240, jc="left", sangria=None, keep=False):
    ppr = p._p.get_or_add_pPr()
    if keep:
        ppr.append(OxmlElement("w:keepNext"))
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), str(antes))
    sp.set(qn("w:after"), str(despues))
    sp.set(qn("w:line"), str(linea))
    sp.set(qn("w:lineRule"), "auto")
    ppr.append(sp)
    if sangria:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLine"), str(sangria))
        ppr.append(ind)
    j = OxmlElement("w:jc")
    j.set(qn("w:val"), jc)
    ppr.append(j)


def hipervinculo(doc, p, url, texto, sz):
    rid = doc.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    s = OxmlElement("w:sz")
    s.set(qn("w:val"), str(sz))
    rpr.append(s)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    p._p.append(h)


def parrafo(doc, texto="", **kw):
    p = doc.add_paragraph()
    espaciado(p, **kw)
    return p


def tabla(doc, encabezados, filas, anchos, sz=20):
    """Tabla con solo líneas horizontales (formato UA); una celda (texto, url) se vuelve hipervínculo."""
    t = doc.add_table(rows=0, cols=len(encabezados))
    tblpr = t._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(sum(anchos)))
    w.set(qn("w:type"), "dxa")
    tblpr.append(w)
    bordes = OxmlElement("w:tblBorders")
    for lado, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"),
                      ("insideH", "nil"), ("insideV", "nil")):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), val)
        if val == "single":
            b.set(qn("w:sz"), "6")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
        bordes.append(b)
    tblpr.append(bordes)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = t._tbl.tblGrid
    for gc, a in zip(grid.findall(qn("w:gridCol")), anchos):
        gc.set(qn("w:w"), str(a))

    def fila(valores, cabecera=False):
        tr = t.add_row()
        trpr = tr._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if cabecera:
            trpr.append(OxmlElement("w:tblHeader"))
        for k, (c, v) in enumerate(zip(tr.cells, valores)):
            c.width = Twips(anchos[k])
            if cabecera:
                tcpr = c._tc.get_or_add_tcPr()
                tb = OxmlElement("w:tcBorders")
                bot = OxmlElement("w:bottom")
                bot.set(qn("w:val"), "single")
                bot.set(qn("w:sz"), "6")
                bot.set(qn("w:space"), "0")
                bot.set(qn("w:color"), "000000")
                tb.append(bot)
                tcpr.append(tb)
            p = c.paragraphs[0]
            espaciado(p, antes=20, despues=20)
            if isinstance(v, tuple):  # (texto, url)
                hipervinculo(doc, p, v[1], v[0], sz) if v[1] else run(p, v[0], sz=sz)
            else:
                run(p, str(v), b=cabecera, sz=sz)

    fila(encabezados, cabecera=True)
    for f in filas:
        fila(f)
    return t


def limpiar_anexos_previos(doc):
    body = doc.element.body
    hijos = list(body)
    for idx, el in enumerate(hijos):
        texto = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip() if el.tag == qn("w:p") else ""
        if texto == "Anexos":
            # el párrafo anterior contiene el salto de sección insertado previamente
            previo = hijos[idx - 1]
            sect_previo = previo.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            final = body.find(qn("w:sectPr"))
            for borrar in hijos[idx:]:
                if borrar is not final:
                    body.remove(borrar)
            if sect_previo is not None:
                body.remove(final)
                body.append(sect_previo)
                body.remove(previo)
            return


def construir_docx(sel, total_incluidos):
    doc = docx.Document(ARTICULO)
    limpiar_anexos_previos(doc)
    body = doc.element.body
    sect_final = body.find(qn("w:sectPr"))
    # los márgenes originales vienen como "1440.0000000000002"; Word exige enteros
    for mar in body.iter(qn("w:pgMar")):
        for k, v in mar.attrib.items():
            mar.set(k, str(round(float(v))))

    # Salto de sección: la sección principal conserva encabezados y numeración
    salto = doc.add_paragraph()
    salto._p.get_or_add_pPr().append(copy.deepcopy(sect_final))

    # Sección de anexos: horizontal, sin numeración de página (guía UA)
    for hijo in list(sect_final):
        if hijo.tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:titlePg"), qn("w:pgNumType")):
            sect_final.remove(hijo)
    seccion = doc.sections[-1]
    seccion.orientation = WD_ORIENT.LANDSCAPE
    seccion.page_width, seccion.page_height = Twips(16838), Twips(11906)
    seccion.header.is_linked_to_previous = False
    seccion.footer.is_linked_to_previous = False
    ancho = 16838 - 2 * 1440

    p = parrafo(doc, despues=120, linea=360)
    run(p, "Anexos", b=True, sz=24)

    # Anexo 1
    p = parrafo(doc, antes=120, keep=True)
    run(p, "Anexo 1", b=True, sz=22)
    p = parrafo(doc, despues=60, keep=True)
    run(p, f"Matriz de los {N_SELECCION} artículos seleccionados", i=True, sz=22)
    anchos = [500, 2000, 5000, 1300, 1450, 1900, 0]
    anchos[-1] = ancho - sum(anchos[:-1])
    filas = [[str(e["n"]), e["autor_anio"], e["titulo"], e["base"], e["pais"],
              etiqueta(TIPO, e["tipo_estudio"]), (e["doi"] or "Ver enlace", e["url_doi"] or e["url_texto"])]
             for e in sel]
    tabla(doc, ["N°", "Autor(es) y año", "Título", "Base de datos", "País", "Tipo de estudio", "DOI"],
          filas, anchos)
    p = parrafo(doc, antes=60, despues=120)
    run(p, "Nota. ", i=True, sz=20)
    run(p, (f"Los {N_SELECCION} artículos se seleccionaron de los {total_incluidos} estudios incluidos en el flujo "
            "PRISMA, priorizando los citados en la Introducción, el mayor puntaje de calidad metodológica, la "
            "evidencia latinoamericana y el número de citas por año. La matriz de extracción completa (muestra, "
            "chatbot, madurez, hallazgos, barreras y limitaciones) se presenta en el archivo "
            "Anexo_50_articulos_seleccionados.xlsx. Elaboración propia."), sz=20)

    # Anexo 2
    p = parrafo(doc, antes=240, keep=True)
    p.paragraph_format.page_break_before = True
    run(p, "Anexo 2", b=True, sz=22)
    p = parrafo(doc, despues=60, keep=True)
    run(p, "Evidencia de recuperación y evaluación de calidad de los artículos seleccionados", i=True, sz=22)
    anchos = [500, 700, 2000, 1500, 0, 520, 520, 520, 520, 520, 950]
    anchos[4] = ancho - sum(anchos) + anchos[4]
    filas = []
    for e in sel:
        url = e["url_texto"] or e["url_doi"]
        filas.append([str(e["n"]), e["id"], e["autor_anio"], e["base"], (url, url)]
                     + [e[q] for q in ("q1", "q2", "q3", "q4", "q5")] + [f"{e['puntaje']}/5"])
    tabla(doc, ["N°", "ID", "Autor(es) y año", "Base de datos", "Enlace al texto completo consultado",
                "C1", "C2", "C3", "C4", "C5", "Puntaje"], filas, anchos)
    p = parrafo(doc, antes=60, despues=120)
    run(p, "Nota. ", i=True, sz=20)
    run(p, ("ID = identificador del registro en las hojas Para_cribado_1171 y Texto_completo del libro "
            "PRISMA_identificacion_consolidado.xlsx. Todos los artículos superaron el cribado por título y "
            "resumen y fueron incluidos tras la lectura del texto completo. C1-C5 = criterios de calidad "
            "metodológica (1 = cumple, 0 = no cumple); se excluyeron los estudios con puntaje menor a 4. "
            "Elaboración propia."), sz=20)

    # Anexo 3
    p = parrafo(doc, antes=240, keep=True)
    run(p, "Anexo 3", b=True, sz=22)
    p = parrafo(doc, despues=60, keep=True)
    run(p, "Presentación en Excel de los artículos seleccionados", i=True, sz=22)
    p = parrafo(doc, linea=360, jc="both", sangria=720)
    run(p, ("El archivo Anexo_50_articulos_seleccionados.xlsx acompaña al presente artículo y contiene cuatro "
            "hojas: Presentacion, con el flujo PRISMA y el criterio de selección; Articulos_50, con la matriz de "
            "extracción de cada artículo y su referencia en formato IEEE; Evidencia, con la trazabilidad de cada "
            "artículo hasta su registro de cribado, el enlace al texto completo consultado y los criterios de "
            "calidad; y Resumen, con la distribución de los artículos por sector, tipo de estudio, año, base de "
            "datos y región."), sz=24)

    doc.save(ARTICULO)


def main():
    estudios, resumen = cargar()
    sel = seleccionar(estudios)
    construir_excel(sel, resumen, len(estudios))
    construir_docx(sel, len(estudios))
    print(f"{len(sel)} artículos seleccionados de {len(estudios)} incluidos")
    print("Sector:", Counter(etiqueta(SECTOR, e["sector"]) for e in sel).most_common())
    print(f"Excel: {SALIDA_XLSX.relative_to(ROOT)}")
    print(f"Word:  {ARTICULO.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
